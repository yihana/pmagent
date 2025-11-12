#  D:\workspace\pm-agent\server\workflow\agents\schedule_agent\outputs\change_mgmt.py
import openpyxl
from pathlib import Path
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment


class ChangeManagementGenerator:
    """변경관리 대장 Excel 생성기"""
    
    @staticmethod
    def generate(project_id: str, output_path: Path) -> str:
        """
        변경관리 대장 Excel 생성
        
        Args:
            project_id: 프로젝트 ID
            output_path: 저장할 파일 경로
            
        Returns:
            str: 생성된 파일의 절대 경로
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "변경관리"
        
        # 헤더 (Image 6 기준)
        headers = [
            "번호", "변경요청일자", "변경요청자", "변경관리담당자", "변경요청명", 
            "변경요청내역", "변경영향(DD일)", "투입인력명(MD)", "장애/개선 명칭",
            "분기/정기(백로그)", "예상 리스크", "합의영향", "변경이행담당자", 
            "변경이행완료", "관련자료"
        ]
        ws.append(headers)
        
        # 헤더 스타일링
        header_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True, size=10)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # 샘플 데이터 (실제로는 변경 이력 추적 시스템과 연동)
        sample_data = [
            [
                1,
                "2023.06.15",
                "○○○ 파트너",
                "김무허사",
                "기능개선",
                "Test 환경을 추가한 이지 변경으로 Test 환경가 작성한 노하우",
                "N일",
                "추운 음주 우기 = M/D",
                "적격사무요청",
                "O초 백로그",
                "기능 능등 자출",
                "",
                "기능 능등 자출",
                "2023.07.01",
                "요구사유 증감사"
            ]
        ]
        
        for row in sample_data:
            ws.append(row)
        
        # 열 너비 조정
        column_widths = {
            'A': 8,   # 번호
            'B': 12,  # 변경요청일자
            'C': 12,  # 변경요청자
            'D': 12,  # 변경관리담당자
            'E': 15,  # 변경요청명
            'F': 40,  # 변경요청내역
            'G': 12,  # 변경영향
            'H': 15,  # 투입인력명
            'I': 15,  # 장애/개선
            'J': 15,  # 분기/정기
            'K': 20,  # 예상 리스크
            'L': 15,  # 합의영향
            'M': 15,  # 변경이행담당자
            'N': 12,  # 변경이행완료
            'O': 15,  # 관련자료
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # 행 높이 조정
        ws.row_dimensions[1].height = 30
        
        # 모든 셀 줄바꿈 활성화
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        
        # 파일 저장
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        
        return str(output_path.resolve())